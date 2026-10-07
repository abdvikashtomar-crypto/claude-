"""Turn a guest post Markdown file into a clean Word document for editors.

Links become "anchor (URL)" so editors can add them by hand, and images become
"[IMAGE: file]" placeholders because most sites want images attached separately.
Usage: python3 backlinks/scripts/md-to-docx.py input.md output.docx
"""
import re, sys
from docx import Document
from docx.shared import Pt

src, out = sys.argv[1], sys.argv[2]
text = re.sub(r'<!--.*?-->', '', open(src).read(), flags=re.S).strip()
doc = Document()
style = doc.styles['Normal']; style.font.name = 'Calibri'; style.font.size = Pt(11)

def add_runs(par, t):
    t = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1 (\2)', t)
    for part in re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)', t):
        if part.startswith('**') and part.endswith('**'):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            par.add_run(part[1:-1]).italic = True
        elif part:
            par.add_run(part)

lines = text.split('\n'); i = 0
while i < len(lines):
    ln = lines[i].rstrip()
    if not ln or ln == '---':
        i += 1; continue
    if ln.startswith('# '):
        doc.add_heading(ln[2:], level=1)
    elif ln.startswith('## '):
        doc.add_heading(ln[3:], level=2)
    elif m := re.match(r'!\[([^\]]*)\]\(([^)]+)\)', ln):
        p = doc.add_paragraph(); p.add_run(f'[IMAGE: {m.group(2)} (attached). Alt text: {m.group(1)}]').italic = True
    elif re.match(r'\d+\. ', ln):
        add_runs(doc.add_paragraph(style='List Number'), re.sub(r'^\d+\. ', '', ln))
    elif ln.startswith('- '):
        add_runs(doc.add_paragraph(style='List Bullet'), ln[2:])
    else:
        add_runs(doc.add_paragraph(), ln)
    i += 1
doc.save(out)
print('wrote', out)
